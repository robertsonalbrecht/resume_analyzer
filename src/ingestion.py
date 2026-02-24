"""Resume file ingestion: PDF and DOCX text extraction."""

import fitz  # PyMuPDF
from docx import Document


PAGE_SENTINEL = "\n--- PAGE BREAK ---\n"


class IngestionError(Exception):
    """Raised when a resume file cannot be ingested."""


def extract_text_from_pdf(file_path: str) -> str:
    """Extract plain text from a PDF file, joining pages with a sentinel."""
    doc = fitz.open(file_path)
    pages = [page.get_text("text") for page in doc]
    doc.close()
    return PAGE_SENTINEL.join(pages)


def extract_text_from_docx(file_path: str) -> str:
    """Extract plain text from a DOCX file (paragraphs + table cells)."""
    document = Document(file_path)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)

    return "\n".join(parts)


def load_resume(file_path: str) -> str:
    """Load a resume from a PDF or DOCX file and return its text.

    Raises IngestionError if the file type is unsupported or the file is empty.
    """
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_path)
    elif lower.endswith(".docx"):
        text = extract_text_from_docx(file_path)
    else:
        raise IngestionError(f"Unsupported file type: {file_path!r}. Only PDF and DOCX are supported.")

    if not text or not text.strip():
        raise IngestionError(f"No text extracted from {file_path!r}.")

    if len(text.strip()) < 200:
        raise IngestionError(
            f"Extracted text from {file_path!r} is too short ({len(text.strip())} chars). "
            "This is likely a scanned or image-only PDF."
        )

    return text
